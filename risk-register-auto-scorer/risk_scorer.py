#!/usr/bin/env python3
"""
Risk Register Auto-Scorer — by Muhammad Bilal
==============================================
Reads risk-register.csv, computes exposure = probability × impact,
assigns RAG bands, draws a 5×5 heatmap chart, and produces a styled
HTML report + a scored Excel register.

100% LOCAL: no APIs, no network, no cloud, no credentials.
Run:  python risk_scorer.py            (uses risk-register.csv in this folder)
      python risk_scorer.py --input my_risks.csv --outdir ./out
"""
from __future__ import annotations

import argparse
import base64
import csv
import logging
import sys
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # headless — safe in scheduled tasks
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- CONFIG ----
CONFIG = {
    "input_csv": "risk-register.csv",
    "output_dir": "output",
    "heatmap_file": "risk-heatmap.png",
    "report_title": "Risk Register Report",
    "bands": {  # (min, max, label, hex)
        "Red": (15, 25, "Red", "#DC2626"),
        "Amber": (8, 14, "Amber", "#F59E0B"),
        "Green": (1, 7, "Green", "#16A34A"),
    },
    "top_n": 5,  # how many top risks to feature
}

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
log = logging.getLogger("risk-scorer")


def argb(hexcol: str, alpha: str = "FF") -> str:
    """Convert '#RRGGBB' → 'AARRGGBB' (openpyxl requires aRGB)."""
    return alpha + hexcol.lstrip("#")


def band_for(exposure: int) -> tuple[str, str]:
    """Return (band label, hex colour) for an exposure score."""
    for label, (lo, hi, _, hexcol) in CONFIG["bands"].items():
        if lo <= exposure <= hi:
            return label, hexcol
    raise ValueError(f"Exposure {exposure} outside 1–25 range")


def load_risks(path: Path) -> list[dict]:
    if not path.exists():
        raise FileNotFoundError(f"Risk register not found: {path}")
    risks: list[dict] = []
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        required = {"risk_id", "risk", "category", "owner",
                    "probability", "impact", "mitigation"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"risk-register.csv missing columns: {sorted(missing)}")
        for i, row in enumerate(reader, start=2):
            try:
                prob = int(row["probability"])
                impact = int(row["impact"])
            except ValueError as exc:
                raise ValueError(
                    f"Row {i} (risk {row.get('risk_id')}): probability/impact "
                    f"must be integers 1–5") from exc
            if not (1 <= prob <= 5 and 1 <= impact <= 5):
                raise ValueError(
                    f"Row {i} (risk {row.get('risk_id')}): probability/impact "
                    f"must be 1–5, got {prob}/{impact}")
            exposure = prob * impact
            band, colour = band_for(exposure)
            risks.append({
                "risk_id": row["risk_id"].strip(),
                "risk": row["risk"].strip(),
                "category": row["category"].strip(),
                "owner": row["owner"].strip(),
                "probability": prob, "impact": impact,
                "exposure": exposure, "band": band, "colour": colour,
                "mitigation": row["mitigation"].strip(),
            })
    log.info("Loaded %d risks from %s", len(risks), path)
    return risks


def draw_heatmap(risks: list[dict], out: Path) -> Path:
    """5×5 probability × impact heatmap; cells show risk IDs in that cell."""
    grid: dict[tuple[int, int], list[str]] = {}
    for r in risks:
        grid.setdefault((r["impact"], r["probability"]), []).append(r["risk_id"])
    fig, ax = plt.subplots(figsize=(8.4, 6.8))
    data = np.zeros((5, 5))
    for impact in range(1, 6):
        for prob in range(1, 6):
            data[5 - impact, prob - 1] = prob * impact
    cmap = plt.cm.YlOrRd
    ax.imshow(data, cmap=cmap, vmin=0, vmax=25, aspect="auto")
    for impact in range(1, 6):
        for prob in range(1, 6):
            ids = grid.get((impact, prob), [])
            label = "\n".join(ids) if ids else str(prob * impact)
            ax.text(prob - 1, 5 - impact, label, ha="center", va="center",
                    fontsize=9 if ids else 11,
                    fontweight="bold" if ids else "normal",
                    color="white" if prob * impact >= 15 else "#1f2937")
    ax.set_xticks(range(5))
    ax.set_xticklabels([f"P{i}" for i in range(1, 6)])
    ax.set_yticks(range(5))
    ax.set_yticklabels([f"I{i}" for i in range(5, 0, -1)])
    ax.set_xlabel("Probability →", fontweight="bold")
    ax.set_ylabel("Impact →", fontweight="bold")
    ax.set_title("Risk Heatmap — Probability × Impact\n(Red ≥ 15 · Amber 8–14 · Green < 8)",
                 fontweight="bold", pad=12)
    fig.tight_layout()
    fig.savefig(out, dpi=150)
    plt.close(fig)
    return out


def write_xlsx(risks: list[dict], today: date, out: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scored Register"
    ws.merge_cells("A1:I1")
    ws["A1"] = f"{CONFIG['report_title']} — {today} · by Muhammad Bilal"
    ws["A1"].font = Font(bold=True, size=14, color="0E4C81")
    headers = ["Risk ID", "Risk", "Category", "Owner", "Probability (1–5)",
               "Impact (1–5)", "Exposure (P×I)", "Band", "Mitigation"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF0E4C81")
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ordered = sorted(risks, key=lambda r: (-r["exposure"], r["risk_id"]))
    for r, risk in enumerate(ordered, 4):
        values = [risk["risk_id"], risk["risk"], risk["category"],
                  risk["owner"], risk["probability"], risk["impact"],
                  risk["exposure"], risk["band"], risk["mitigation"]]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c in (7, 8):
                cell.fill = PatternFill("solid",
                                        fgColor=argb(risk["colour"], "33"))
                cell.font = Font(bold=True)
    for c, w in enumerate([10, 52, 16, 16, 13, 12, 13, 9, 52], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:I{3 + len(risks)}"
    wb.save(out)
    return out


def write_html(risks: list[dict], today: date, heatmap: Path, out: Path) -> Path:
    ordered = sorted(risks, key=lambda r: (-r["exposure"], r["risk_id"]))
    counts = {b: sum(1 for r in risks if r["band"] == b)
              for b in CONFIG["bands"]}
    img_b64 = base64.b64encode(heatmap.read_bytes()).decode()
    top_rows = ""
    for r in ordered[:CONFIG["top_n"]]:
        top_rows += (
            f'<tr><td><span class="badge" style="background:{r["colour"]}22;'
            f'color:{r["colour"]}">{r["band"]}</span></td>'
            f'<td>{r["risk_id"]}</td><td>{r["risk"]}</td>'
            f'<td>{r["category"]}</td><td>{r["owner"]}</td>'
            f'<td>{r["probability"]} × {r["impact"]}</td>'
            f'<td><b>{r["exposure"]}</b></td><td>{r["mitigation"]}</td></tr>\n')
    all_rows = ""
    for r in ordered:
        all_rows += (
            f'<tr><td><span class="badge" style="background:{r["colour"]}22;'
            f'color:{r["colour"]}">{r["band"]}</span></td>'
            f'<td>{r["risk_id"]}</td><td>{r["risk"]}</td>'
            f'<td>{r["category"]}</td><td>{r["owner"]}</td>'
            f'<td>{r["probability"]} × {r["impact"]}</td>'
            f'<td><b>{r["exposure"]}</b></td></tr>\n')
    html = f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8">
<title>{CONFIG['report_title']} — {today}</title>
<style>
body{{font-family:system-ui,-apple-system,Segoe UI,Roboto,sans-serif;margin:32px;color:#1f2937;background:#f8fafc}}
h1{{margin:0 0 4px}} .sub{{color:#6b7280;margin-bottom:20px}}
.summary{{display:flex;gap:12px;margin-bottom:20px}}
.card{{background:#fff;border-radius:10px;padding:14px 18px;box-shadow:0 1px 3px rgba(0,0,0,.08);min-width:110px}}
.card b{{font-size:26px;display:block}} .card span{{color:#6b7280;font-size:13px}}
table{{border-collapse:collapse;width:100%;background:#fff;border-radius:10px;overflow:hidden;margin-bottom:24px}}
th,td{{padding:10px 12px;border-bottom:1px solid #e5e7eb;text-align:left;font-size:14px}}
th{{background:#0E4C81;color:#fff}} .badge{{padding:3px 10px;border-radius:999px;font-weight:700;font-size:12px}}
.heatmap{{background:#fff;border-radius:10px;padding:16px;text-align:center;box-shadow:0 1px 3px rgba(0,0,0,.08);margin-bottom:24px}}
.heatmap img{{max-width:640px;width:100%}}
footer{{color:#6b7280;font-size:12px}}
</style></head><body>
<h1>{CONFIG['report_title']}</h1>
<p class="sub">Generated {today.isoformat()} · Risk Register Auto-Scorer · by Muhammad Bilal · 100% local, no external data</p>
<div class="summary">
<div class="card"><b style="color:#DC2626">{counts['Red']}</b><span>Red (≥15)</span></div>
<div class="card"><b style="color:#F59E0B">{counts['Amber']}</b><span>Amber (8–14)</span></div>
<div class="card"><b style="color:#16A34A">{counts['Green']}</b><span>Green (&lt;8)</span></div>
<div class="card"><b>{len(risks)}</b><span>Total risks</span></div>
</div>
<h2>Top {CONFIG['top_n']} Risks</h2>
<table><thead><tr><th>Band</th><th>Risk ID</th><th>Risk</th><th>Category</th><th>Owner</th>
<th>P × I</th><th>Exposure</th><th>Mitigation</th></tr></thead><tbody>{top_rows}</tbody></table>
<div class="heatmap"><h2 style="text-align:left">Risk Heatmap</h2>
<img src="data:image/png;base64,{img_b64}" alt="5x5 risk heatmap"></div>
<h2>Full Scored Register</h2>
<table><thead><tr><th>Band</th><th>Risk ID</th><th>Risk</th><th>Category</th><th>Owner</th>
<th>P × I</th><th>Exposure</th></tr></thead><tbody>{all_rows}</tbody></table>
<footer>Exposure = Probability (1–5) × Impact (1–5). RAG bands: Red ≥ 15 · Amber 8–14 · Green &lt; 8.
Scores come from the register only — review Red risks in the next steering meeting.</footer>
</body></html>"""
    out.write_text(html, encoding="utf-8")
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Risk Register Auto-Scorer")
    ap.add_argument("--input", default=CONFIG["input_csv"])
    ap.add_argument("--outdir", default=CONFIG["output_dir"])
    args = ap.parse_args(argv)

    base = Path(__file__).resolve().parent
    input_path = Path(args.input) if Path(args.input).is_absolute() \
        else base / args.input
    outdir = Path(args.outdir) if Path(args.outdir).is_absolute() \
        else base / args.outdir
    outdir.mkdir(parents=True, exist_ok=True)
    today = date.today()

    try:
        risks = load_risks(input_path)
        stamp = today.isoformat()
        heatmap = draw_heatmap(risks, outdir / CONFIG["heatmap_file"])
        html = write_html(risks, today, heatmap, outdir / f"risk-report-{stamp}.html")
        xlsx = write_xlsx(risks, today, outdir / f"scored-register-{stamp}.xlsx")
        counts = {b: sum(1 for r in risks if r["band"] == b)
                  for b in CONFIG["bands"]}
        log.info("DONE — %d risks (Red=%d Amber=%d Green=%d)",
                 len(risks), counts["Red"], counts["Amber"], counts["Green"])
        log.info("Heatmap: %s", heatmap)
        log.info("HTML: %s", html)
        log.info("XLSX: %s", xlsx)
        return 0
    except (FileNotFoundError, ValueError) as exc:
        log.error("FAILED: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
