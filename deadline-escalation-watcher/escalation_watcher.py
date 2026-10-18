#!/usr/bin/env python3
"""
Deadline Escalation Watcher — by Muhammad Bilal
================================================
Reads tasks.csv, flags overdue / due-soon tasks, assigns severity tiers,
and produces a styled HTML report, an Excel report, and an idempotent
append-only escalation log.

100% LOCAL: no APIs, no network, no cloud, no credentials.
Run:  python escalation_watcher.py            (uses tasks.csv in this folder)
      python escalation_watcher.py --input my_tasks.csv --outdir ./out
"""
from __future__ import annotations

import argparse
import csv
import logging
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- CONFIG ----
# All behaviour knobs live here — edit without touching the logic below.
CONFIG = {
    "input_csv": "tasks.csv",            # task register (CSV)
    "output_dir": "output",              # where reports + log are written
    "log_file": "escalation-log.csv",    # append-only log (dedupe: task_id + date)
    "overdue_critical_days": 7,          # overdue > this => Critical
    "due_soon_hours": 48,                # due within this => Medium
    "due_this_week_days": 7,              # due within this => Watch
    "ignore_statuses": {"done", "completed", "cancelled", "closed"},
    "report_title": "Deadline Escalation Report",
}

SEVERITY_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Watch": 3}
SEVERITY_STYLE = {
    "Critical": ("#FCE8E8", "#B91C1C"),
    "High": ("#FFF3CD", "#92400E"),
    "Medium": ("#FEF3C7", "#92400E"),
    "Watch": ("#E0F2FE", "#1E40AF"),
}


def argb(hexcol: str, alpha: str = "FF") -> str:
    """Convert '#RRGGBB' → 'AARRGGBB' (openpyxl requires aRGB)."""
    return alpha + hexcol.lstrip("#")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
log = logging.getLogger("escalation-watcher")


# ------------------------------------------------------------------ I/O ----
def load_tasks(path: Path) -> list[dict]:
    """Load and validate the task register."""
    if not path.exists():
        raise FileNotFoundError(f"Task register not found: {path}")
    tasks: list[dict] = []
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        required = {"task_id", "project", "task_name", "owner", "due_date",
                    "status", "blocks_milestone"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"tasks.csv missing columns: {sorted(missing)}")
        for i, row in enumerate(reader, start=2):
            try:
                due = datetime.strptime(row["due_date"].strip(), "%Y-%m-%d").date()
            except ValueError as exc:
                raise ValueError(
                    f"Row {i} (task {row.get('task_id')}): bad due_date "
                    f"'{row.get('due_date')}' — expected YYYY-MM-DD") from exc
            tasks.append({
                "task_id": row["task_id"].strip(),
                "project": row["project"].strip(),
                "task_name": row["task_name"].strip(),
                "owner": row["owner"].strip(),
                "due_date": due,
                "status": row["status"].strip(),
                "blocks_milestone": row["blocks_milestone"].strip().lower() in
                                    {"yes", "y", "true", "1"},
            })
    log.info("Loaded %d tasks from %s", len(tasks), path)
    return tasks


# ------------------------------------------------------------ CLASSIFY ----
def classify(task: dict, today: date, cfg: dict) -> dict | None:
    """Return the escalation record for a task, or None if no action needed."""
    if task["status"].lower() in cfg["ignore_statuses"]:
        return None
    delta = (task["due_date"] - today).days
    if delta < 0:                                   # overdue
        overdue_days = -delta
        if overdue_days > cfg["overdue_critical_days"] or task["blocks_milestone"]:
            severity, bucket = "Critical", "Overdue"
        else:
            severity, bucket = "High", "Overdue"
        detail = f"Overdue by {overdue_days} day(s)"
    elif delta * 24 <= cfg["due_soon_hours"]:       # due within 48h
        severity, bucket = "Medium", "Due in 48h"
        detail = f"Due in {delta} day(s) — {task['due_date']}"
    elif delta <= cfg["due_this_week_days"]:        # due this week
        severity, bucket = "Watch", "Due this week"
        detail = f"Due in {delta} day(s) — {task['due_date']}"
    else:
        return None
    return {
        "task_id": task["task_id"], "project": task["project"],
        "task_name": task["task_name"], "owner": task["owner"],
        "due_date": task["due_date"].isoformat(), "status": task["status"],
        "blocks_milestone": "Yes" if task["blocks_milestone"] else "No",
        "severity": severity, "bucket": bucket, "detail": detail,
        "days_overdue": max(0, -delta),
    }


# --------------------------------------------------------------- OUTPUTS ----
def write_html(records: list[dict], today: date, out: Path) -> Path:
    """Styled, severity-coloured HTML report."""
    counts = {s: sum(1 for r in records if r["severity"] == s)
              for s in SEVERITY_ORDER}
    rows = ""
    for r in sorted(records, key=lambda x: (SEVERITY_ORDER[x["severity"]],
                                             x["due_date"])):
        bg, fg = SEVERITY_STYLE[r["severity"]]
        rows += (
            f'<tr><td><span class="badge" style="background:{bg};color:{fg}">'
            f'{r["severity"]}</span></td><td>{r["task_id"]}</td>'
            f'<td>{r["project"]}</td><td>{r["task_name"]}</td>'
            f'<td>{r["owner"]}</td><td>{r["due_date"]}</td>'
            f'<td>{r["bucket"]}</td><td>{r["detail"]}</td>'
            f'<td>{"YES" if r["blocks_milestone"] == "Yes" else "—"}</td></tr>\n')
    if not rows:
        rows = ('<tr><td colspan="9" style="text-align:center">'
                'No escalations today — all tasks on track.</td></tr>')
    html = f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8">
<title>{CONFIG['report_title']} — {today}</title>
<style>
body{{font-family:system-ui,-apple-system,Segoe UI,Roboto,sans-serif;margin:32px;color:#1f2937;background:#f8fafc}}
h1{{margin:0 0 4px}} .sub{{color:#6b7280;margin-bottom:20px}}
.summary{{display:flex;gap:12px;margin-bottom:20px}}
.card{{background:#fff;border-radius:10px;padding:14px 18px;box-shadow:0 1px 3px rgba(0,0,0,.08);min-width:110px}}
.card b{{font-size:26px;display:block}} .card span{{color:#6b7280;font-size:13px}}
table{{border-collapse:collapse;width:100%;background:#fff;border-radius:10px;overflow:hidden}}
th,td{{padding:10px 12px;border-bottom:1px solid #e5e7eb;text-align:left;font-size:14px}}
th{{background:#0E4C81;color:#fff}} .badge{{padding:3px 10px;border-radius:999px;font-weight:700;font-size:12px}}
footer{{margin-top:16px;color:#6b7280;font-size:12px}}
</style></head><body>
<h1>{CONFIG['report_title']}</h1>
<p class="sub">Generated {today.isoformat()} · Deadline Escalation Watcher · by Muhammad Bilal · 100% local, no external data</p>
<div class="summary">
<div class="card"><b style="color:#B91C1C">{counts['Critical']}</b><span>Critical</span></div>
<div class="card"><b style="color:#92400E">{counts['High']}</b><span>High</span></div>
<div class="card"><b style="color:#92400E">{counts['Medium']}</b><span>Medium</span></div>
<div class="card"><b style="color:#1E40AF">{counts['Watch']}</b><span>Watch</span></div>
<div class="card"><b>{len(records)}</b><span>Total escalations</span></div>
</div>
<table><thead><tr><th>Severity</th><th>Task ID</th><th>Project</th><th>Task</th>
<th>Owner</th><th>Due date</th><th>Bucket</th><th>Detail</th><th>Blocks milestone</th></tr></thead>
<tbody>{rows}</tbody></table>
<footer>Severity rules: Critical = overdue &gt; 7 days OR blocks a milestone · High = overdue ≤ 7 days ·
Medium = due within 48h · Watch = due this week. Completed / cancelled tasks are ignored.</footer>
</body></html>"""
    out.write_text(html, encoding="utf-8")
    return out


def write_xlsx(records: list[dict], today: date, out: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Escalations"
    ws.merge_cells("A1:I1")
    ws["A1"] = f"{CONFIG['report_title']} — {today} · by Muhammad Bilal"
    ws["A1"].font = Font(bold=True, size=14, color="0E4C81")
    headers = ["Severity", "Task ID", "Project", "Task", "Owner", "Due date",
               "Bucket", "Detail", "Blocks milestone"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF0E4C81")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, rec in enumerate(sorted(records,
                                   key=lambda x: (SEVERITY_ORDER[x["severity"]],
                                                  x["due_date"])), 4):
        bg, fg = SEVERITY_STYLE[rec["severity"]]
        values = [rec["severity"], rec["task_id"], rec["project"],
                  rec["task_name"], rec["owner"], rec["due_date"],
                  rec["bucket"], rec["detail"], rec["blocks_milestone"]]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill = PatternFill("solid", fgColor=argb(bg))
            cell.font = Font(color=argb(fg), bold=(c == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for c, w in enumerate([11, 12, 22, 46, 18, 12, 14, 30, 16], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:I{3 + max(len(records), 1)}"
    wb.save(out)
    return out


def append_log(records: list[dict], today: date, log_path: Path) -> int:
    """Append new entries only — idempotent on (task_id, date)."""
    existing: set[tuple[str, str]] = set()
    if log_path.exists():
        with log_path.open(newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                existing.add((row["task_id"], row["report_date"]))
    new_rows = [r for r in records
                if (r["task_id"], today.isoformat()) not in existing]
    if not new_rows:
        log.info("Log already up to date for %s — 0 new entries", today)
        return 0
    write_header = not log_path.exists()
    with log_path.open("a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if write_header:
            w.writerow(["report_date", "task_id", "project", "task_name",
                        "owner", "due_date", "severity", "bucket", "detail"])
        for r in new_rows:
            w.writerow([today.isoformat(), r["task_id"], r["project"],
                        r["task_name"], r["owner"], r["due_date"],
                        r["severity"], r["bucket"], r["detail"]])
    log.info("Appended %d new log entries", len(new_rows))
    return len(new_rows)


# ------------------------------------------------------------------ MAIN ----
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Deadline Escalation Watcher")
    ap.add_argument("--input", default=CONFIG["input_csv"])
    ap.add_argument("--outdir", default=CONFIG["output_dir"])
    args = ap.parse_args(argv)

    base = Path(__file__).resolve().parent
    input_path = Path(args.input) if Path(args.input).is_absolute() \
        else base / args.input
    outdir = Path(args.outdir) if Path(args.outdir).is_absolute() \
        else base / args.outdir
    outdir.mkdir(parents=True, exist_ok=True)
    today = date.today()

    try:
        tasks = load_tasks(input_path)
        records = [r for t in tasks if (r := classify(t, today, CONFIG))]
        stamp = today.isoformat()
        html = write_html(records, today, outdir / f"escalation-report-{stamp}.html")
        xlsx = write_xlsx(records, today, outdir / f"escalation-report-{stamp}.xlsx")
        added = append_log(records, today, outdir / CONFIG["log_file"])
        by_sev = {s: sum(1 for r in records if r["severity"] == s)
                  for s in SEVERITY_ORDER}
        log.info("DONE — %d escalations (Critical=%d High=%d Medium=%d Watch=%d)",
                 len(records), by_sev["Critical"], by_sev["High"],
                 by_sev["Medium"], by_sev["Watch"])
        log.info("HTML: %s", html)
        log.info("XLSX: %s", xlsx)
        log.info("Log: %s (+%d new)", outdir / CONFIG["log_file"], added)
        return 0
    except (FileNotFoundError, ValueError) as exc:
        log.error("FAILED: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
